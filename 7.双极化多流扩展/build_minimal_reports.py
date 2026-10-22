"""Build matched Markdown and DOCX explanations from audited minimal-run CSVs."""
from pathlib import Path
import json
import math
import re
from PIL import Image, ImageDraw, ImageFont
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'results/minimal_all_methods_20261022'
TITLE = '单流单极化预编码扩展为双流双极化的最小验证报告'
STEM = '双流双极化全方法最小验证与扩展说明'
PAIRS = pd.read_csv(OUT / 'paired_summary.csv')
AUDIT = json.loads((OUT / 'complete_audit.json').read_text(encoding='utf-8'))
assert AUDIT['status'] == 'passed'
BLOCKS = []


def para(text): BLOCKS.append(('p', text))
def chapter(title): BLOCKS.append(('chapter', title))
def table(title, headers, rows, widths=None): BLOCKS.append(('table', (title, headers, rows, widths)))
def equation(latex, parts): BLOCKS.append(('eq', (latex, parts)))


def pipeline():
    assets = OUT / 'assets';assets.mkdir(exist_ok=True)
    im = Image.new('RGB', (1600, 530), 'white');d = ImageDraw.Draw(im)
    font = ImageFont.truetype('C:/Windows/Fonts/simsun.ttc', 31)
    small = ImageFont.truetype('C:/Windows/Fonts/simsun.ttc', 25)
    boxes = {'y': (20, 190, 270, 330), 'msg': (340, 40, 700, 190), 'f': (780, 40, 1110, 190),
             'ue': (340, 340, 700, 490), 'cpu': (780, 340, 1110, 490), 'w': (1200, 190, 1580, 330)}
    labels = {'y': 'UE观测信道\nY', 'msg': '频率 方向索引\n或坐标消息', 'f': '空间波束\nF_b',
              'ue': 'UE观测及F_b\n量化系数C_b', 'cpu': '几何模型及F_b\nCPU系数C_b', 'w': '空间波束与系数组合\nW_b = F_b C_b'}
    for key, box in boxes.items():
        d.rounded_rectangle(box, 14, fill='#EDF2F6', outline='#36566F', width=3)
        lines = labels[key].split('\n');cx=(box[0]+box[2])/2;cy=(box[1]+box[3])/2
        for j, line in enumerate(lines):
            bb=d.textbbox((0,0),line,font=font);d.text((cx-(bb[2]-bb[0])/2,cy-40+j*47),line,font=font,fill='black')
    def arrow(points):
        d.line(points,fill='#36566F',width=4);x,y=points[-1];px,py=points[-2];a=math.atan2(y-py,x-px)
        d.polygon([(x,y),(x-17*math.cos(a-.45),y-17*math.sin(a-.45)),(x-17*math.cos(a+.45),y-17*math.sin(a+.45))],fill='#36566F')
    arrow([(270,225),(305,225),(305,115),(340,115)])
    arrow([(270,295),(305,295),(305,415),(340,415)])
    arrow([(700,115),(780,115)]);arrow([(1110,115),(1150,115),(1150,225),(1200,225)])
    arrow([(945,190),(945,340)]);arrow([(1110,415),(1150,415),(1150,295),(1200,295)])
    arrow([(700,415),(740,415),(740,265),(1200,265)])
    d.text((110,95),'公共空间信息',font=small,fill='#36566F')
    d.text((30,395),'UE系数链路',font=small,fill='#36566F')
    im.save(assets/'扩展处理流程.png')


def build_content():
    para('本报告说明如何把原低反馈预编码仿真中的单流单极化结构，扩展为双流双极化结构，并用原有全部方法完成最小可行性验证。原方法包括SVD、Type-I、Type-II、HAQSE R3/R5、DFT R3/R5和POS R3/R5，另保留注水容量参考。')
    para('验证结果支持这些方法在当前受控信道模型下形成两个有效接收模式。32个固定位置、两种阵列和两种信道下，共4096条验证记录均通过数值检查；其中1792条可实施双流记录均保持接收秩2。所有UE链路在本次32个位置上均获得正的双流增量；ULA混合信道的DFT CPU链路存在少数负增量位置。可行性成立不代表每个位置都应固定发送两流。')
    para('本次目标是补齐方法并验证实现链路。仅固定B为8、观测SNR为30 dB，不进行SNR、AP数、量化精度或极化失配扫描，也不根据这32个位置给出算法排名。')
    table('表1 本次验证范围', ['项目','设置'], [
        ['样本','原冻结400位置库中的前32个位置  所有方法共同使用'],
        ['独立流程预检查','原库第399和400个位置  与32个验证位置不重叠'],
        ['阵列与信道','ULA64及UPA32×32  LOS及六散射体混合信道'],
        ['发送配置','同一双极化硬件分别发送1流和2流'],
        ['功率与带宽','总功率1  可实施方法每AP功率1/8  带宽20 MHz'],
        ['极化条件','标称已知极化方向  无新增姿态或端口相位失配']
    ],[3.4,13.0])
    para('需要区分两个概念：单流是发送的数据流数为1，并不表示联合多AP信道只能有一个空间模式；双极化描述每个空间阵元的两个极化端口，也不等同于自动获得两个无干扰数据流。')

    chapter('一 从向量预编码扩展为矩阵预编码')
    para('原模型中，每个AP有N个单极化空间阵元，UE有4个接收空间位置，单AP信道为4×N，预编码为N×1向量。原方案利用一个空间波束和一个AP间复系数承载单流。')
    equation(r'\mathbf w_b=\mathbf f_b c_b', [('sub','w','b'),' = ',('sub','f','b'),' ',('sub','c','b')])
    para('扩展时保留原空间坐标、阵元间距和AP朝向。每个空间阵元改为两个±45°极化端口，UE为4个空间位置乘2个极化端口，共8个接收端口。单AP信道变为8×2N，双流预编码变为2N×2矩阵。')
    table('表2 每AP矩阵维度变化', ['场景','原信道','扩展信道','双流预编码'], [
        ['ULA64','4×64','8×128','128×2'],['UPA32×32','4×1024','8×2048','2048×2']
    ],[4.5,3.5,4.5,3.9])
    equation(r'\mathbf W_b=(\mathbf I_2\otimes\mathbf f_b)\mathbf C_b=\mathbf F_b\mathbf C_b', [('sub','W','b'),' = (',('sub','I','2'),' ⊗ ',('sub','f','b'),') ',('sub','C','b'),' = ',('sub','F','b'),' ',('sub','C','b')])
    para('F_b为2N×2，给两个极化端口配置相同的公共空间波束；C_b为2×Ns，其中Ns为发送流数。单流时C_b是2×1向量，双流时为2×2矩阵。它决定两个极化端口的配比、跨AP相对相位及各流的线性组合，因此不能直接把旧的一个共相位复制给两个流。')
    para('真信道保留逐阵元球面传播距离，LOS使用横向投影Jones矩阵描述极化耦合；六条散射路径分别加入极化耦合矩阵，设置XPR为10 dB，再按原K为10 dB的Rician规则合成。极化按中心射线计算，并归一路径极化矩阵范数。这是受控双极化模型，尚不是完整3GPP或实测信道。')
    BLOCKS.append(('image', ('assets/扩展处理流程.png','图1 公共空间波束与两条系数设计链路')))

    chapter('二 原有方法分别怎样扩展')
    para('本次保留原方法获取空间波束的区别，并把后续的AP单流系数统一扩展为双极化多流矩阵。HAQSE、DFT和POS分别运行CPU与UE两条链路；Type-I和Type-II按原仓库定义继续使用UE观测CSI，作为两类链路共享的基线。')
    table('表3 方法与空间消息', ['方法','UE反馈的空间信息','CPU恢复空间波束'], [
        ['HAQSE R3/R5','参考AP的量化空间频率 每维14 bit','由频率定位 参考AP保留直接频率波束 其余AP由位置推断'],
        ['DFT R3/R5','参考AP的DFT波束索引','由索引定位 参考AP保留直接索引波束 其余AP由位置推断'],
        ['POS R3/R5','连续HAQSE方向在UE定位后的量化坐标','全部AP波束由解码坐标生成 包括参考AP'],
        ['Type-I','每AP一个DFT空间波束索引','直接恢复每AP空间波束 两极化共用'],
        ['Type-II','每AP两个空间索引及复合波束的幅相索引','恢复每AP双波束复合结果 两极化共用'],
        ['SVD等功率参考','真实完整信道 不设有限反馈预算','前Ns个真实右奇异向量 每流功率1/Ns'],
        ['注水容量参考','真实完整信道 不设有限反馈预算','最多Ns流注水容量 放松每AP功率约束']
    ],[3.0,6.1,7.3])
    para('双极化参考AP的观测可整理成16个空间快拍。HAQSE按观测能量选4个快拍后调用原估计器。DFT参考方向则由本地观测前两个模式在两种发射极化上的空间能量评分选取；Type-I和Type-II对联合观测前两个模式作相同形式的空间评分。两个发送流数共用所选波束，以便检查在同一子空间内从一流增加到两流的效果。')
    para('Type-II保留原仓库的双波束复合形式。先在原正交DFT分组内选两个能量较强的波束，再由它们的2×2投影协方差主特征向量获得相对幅相。相对幅度量化为3 bit，相对相位量化为2 bit，复合为空间波束后，再乘2×Ns的C_b。此处没有扩大为每极化两个独立波束再接4×Ns系数的更一般结构。')
    para('为检查扩展是否保留原定义，构造单模式且只有一个发射极化有效的退化信道。新Type-I和Type-II解码得到的空间波束、方向bit和内部组合bit均与原实现一致，最大波束差为0。本报告中的Type-I和Type-II指仓库的简化结构扩展，不声称完整标准码本符合性。')

    chapter('三 有效信道 系数编码和功率控制')
    para('把各AP的H_b F_b按列拼接，得到8×2B的有效信道E。UE链路使用带噪观测得到的有效信道，CPU链路用解码位置、公开AP几何和已知标称极化方向重构局部平面LOS有效信道。CPU没有读取真实信道或UE观测CSI。')
    equation(r'\mathbf E=[\mathbf H_1\mathbf F_1,\ldots,\mathbf H_B\mathbf F_B]', ['E = [',('sub','H','1'),' ',('sub','F','1'),', …, ',('sub','H','B'),' ',('sub','F','B'),']'])
    para('以前Ns个右奇异向量初始化C，将每个AP的2×Ns系数块归一到Frobenius范数1/√B。由于F_b的两列正交且单位范数，这保证每AP总功率为1/B，并保证所有AP与所有流合计总功率为1。该逐AP归一是启发式设计，不保证原SVD正交性或每AP功率约束下的全局最优。')
    equation(r'\|\mathbf C_b\|_F^2=1/B,\quad\sum_b\|\mathbf W_b\|_F^2=1', ['‖',('sub','C','b'),'‖² = 1/B,     ',('sum','b=1','B'),' ‖',('sub','W','b'),'‖² = 1'])
    para('UE系数逐复数采用幅度3 bit和相位4 bit，并为每流固定一个不影响速率的全局相位锚。CPU只接收整数索引，解码后再次逐AP归一。原来的跨AP共相位已经包含在系数相位中，不再重复计费。CPU链路直接在CPU计算系数，不发送UE系数消息。')
    equation(r'L_C=2BN_s(q_A+q_P)-N_sq_P', [('sub','L','C'),' = 2B',('sub','N','s'),'(',('sub','q','A'),' + ',('sub','q','P'),') − ',('sub','N','s'),('sub','q','P')])
    para('B为8时，单流系数为108 bit，双流为216 bit。POS坐标统一沿用相位敏感度分配：ULA为24 bit，x/y各12 bit；UPA为36 bit，x/y/z分别13/14/9 bit。R3与R5坐标预算相同。Type-II另外计每AP的两个方向索引和5 bit内部复合系数。')
    bitrows=[]
    for method in ['HAQSE_R3','HAQSE_R5','DFT_R3','DFT_R5','POS_R3','POS_R5','TypeI','TypeII']:
        row=[method.replace('_',' ')]
        for branch in ['CPU','UE']:
            for array in ['ULA','UPA']:
                z=PAIRS[(PAIRS.method==method)&(PAIRS.array==array)&PAIRS.branch.eq(branch if method not in ['TypeI','TypeII'] else 'UE_shared')]
                row.append('UE基线' if method in ['TypeI','TypeII'] and branch=='CPU' else f'{z.dual_bits.iloc[0]:.0f}')
        bitrows.append(row)
    table('表4 B为8时的双流有效载荷 bit', ['方法','CPU ULA','CPU UPA','UE ULA','UE UPA'],bitrows,[4.4,3.0,3.0,3.0,3.0])
    para('流数、码本配置和参考AP集合预先约定。本次不含自适应RI、消息头和导频开销；也不将反馈占用从吞吐量中扣除。旧单极化单流的52/64 bit属于另一种消息格式，不能用来描述本轮双极化UE链路。')

    chapter('四 最小验证怎样运行与判定')
    para('先在第399和400个位置进行独立流程预检查，再用预先固定的前32个位置确认。选择位置时没有依据本轮结果筛选；同一阵列与信道下，所有方法和流数共享真实信道、标准噪声及观测SNR。新端口噪声沿用既有双极化模块的固定种子规则，按当前协作信道平均元素功率缩放。')
    table('表5 验证规模和判据', ['项目','执行内容'], [
        ['方法覆盖','原九方法配置 加注水容量参考'],
        ['每个位置的设计数','6个几何方法乘CPU/UE 加Type-I/II 加SVD/容量 共16种'],
        ['正式记录数','2阵列×2信道×32位置×2流数×16设计 共4096条'],
        ['流程预检查','2个不重叠位置 共256条记录'],
        ['接收双流判据','有效接收矩阵H W的两个奇异值均非零'],
        ['数值检查','整数消息 解码功率 CPU/UE空间消息一致性 速率与容量关系'],
        ['回归检查','32位置全部POS结果与先前400位置存档一致 Type-I波束与既有全方向模块一致']
    ],[3.8,12.6])
    para('主速率使用线性MMSE接收器逐流SINR计算，吞吐量等于20乘以频谱效率，单位Mbps。同时记录联合解码速率和真实信道注水容量，核查MMSE不超过联合解码速率，联合解码速率不超过对应秩的容量参考。单流时MMSE与联合解码结果一致。')
    equation(r'R=20\sum_{s=1}^{N_s}\log_2(1+\mathrm{SINR}_s)', ['R = 20 ',('sum','s=1','Nₛ'),' log₂(1 + ',('sub','SINR','s'),')  Mbps'])
    para('4096条记录通过核查；1792条可实施双流记录全部保持接收秩2，最弱的第二/第一接收模式功率比为0.02731。独立程序重新解码2048条UE系数消息和1024条坐标消息，功率误差低于10⁻¹⁴，无坐标溢出。新增Type-I/II退化回归8项检查通过，预检查中的POS重算差约5.12×10⁻¹³ Mbps，32位置POS与已发布结果差为0。')
    para('接收端评估使用真实有效信道计算MMSE SINR，未另加接收信道估计误差。SVD等功率参考与注水容量参考都使用真实完整信道，并放松每AP功率约束；前者不是一般容量上界。两个参考不属于有限反馈可实施方案。')

    for array in ['ULA','UPA']:
        chapter(('五 ULA最小验证结果' if array=='ULA' else '六 UPA最小验证结果'))
        para('表中为32个位置的平均MMSE吞吐量，单位Mbps。全部使用B为8、观测SNR为30 dB和同一双极化硬件。CPU与UE行分别为六个几何方法的两条链路；Type-I/II为共享UE基线。这里只检查同一方法从一流扩展到两流的行为，不作跨方法排名。')
        rows=[]
        order=[(m,b) for b in ['CPU','UE'] for m in ['HAQSE_R3','HAQSE_R5','DFT_R3','DFT_R5','POS_R3','POS_R5']]+[('TypeI','UE_shared'),('TypeII','UE_shared'),('SVD_equal','ideal'),('CAP_rank_le_N','ideal')]
        for method,branch in order:
            label=method.replace('_',' ')
            if branch in ['CPU','UE']:label+=' '+branch
            elif method=='SVD_equal':label='SVD等功率参考'
            elif method=='CAP_rank_le_N':label='注水容量参考'
            row=[label]
            for channel in ['LOS','MIX']:
                z=PAIRS[(PAIRS.array==array)&(PAIRS.channel==channel)&(PAIRS.method==method)&(PAIRS.branch==branch)].iloc[0]
                row.extend([f'{z.single_Mbps:.2f}',f'{z.dual_Mbps:.2f}'])
            rows.append(row)
        table('表6 '+array+'单流与双流结果' if array=='ULA' else '表7 '+array+'单流与双流结果', ['方法与链路','LOS一流','LOS两流','混合一流','混合两流'],rows,[5.2,2.8,2.8,2.8,2.8])
        r=PAIRS[(PAIRS.array==array)&PAIRS.branch.ne('ideal')]
        para(f'本阵列所有方法在两种信道下的平均双流增量均为正，范围为{r.gain_Mbps.min():.2f}至{r.gain_Mbps.max():.2f} Mbps。实际链路均保持两个非零接收模式，因此HAQSE、DFT、Type-I和Type-II也具备本次模型下的双极化双流可行性。')
        if array=='ULA':
            para('需要保留的失败样本信息是：混合信道中，DFT R3 CPU有28/32个位置获得双流正增量，4个位置降低，最差降低4.98 Mbps；DFT R5 CPU为31/32个位置正增量，1个位置降低，最差降低17.95 Mbps。其余可实施配置本次均为32/32正增量。形成两个非零模式与固定发两流总能提速是不同判据，后续若研究自适应流数，应额外设计RI决策并计入其开销。')
        else:
            para('UPA的所有可实施方法在LOS和混合信道下均为32/32个位置获得正的双流增量。CPU的DFT链路仍比其UE链路存在较大平均差距，与粗方向经定位后参与相干系数设计的敏感性一致；该观察不能替代对定位、极化和多径误差的独立因果分解。')

    chapter('七 当前结论与复现入口')
    para('本轮完成了全部原方法的最小补齐。共同可扩展部分是把AP的单流复系数改为2×Ns矩阵，使极化配比、跨AP相位和流间线性组合在同一低维有效信道上设计；各方法原有的频率、方向、位置或双波束消息仍各自保留。')
    para('当前证据支持实现层面的双流双极化可行性：所有可实施配置满足同总功率和每AP功率约束，全部双流样本保持接收秩2，各方法平均速率相对单流均有增加。DFT CPU的少数负增量位置说明，实际是否发送两流仍与工作点及信道估计质量有关。UE系数来自观测有效信道，CPU系数依赖位置和标称极化模型，两者信息条件不同，应分别解读。')
    para('验证仅限固定部署、32个位置、标称极化、窄带单用户和B为8的一个观测工作点。Type-II采用原复合双波束结构的最小扩展，Type-I/II未做完整标准符合性验证。未覆盖独立部署重复、宽带频率选择性、多用户干扰、移动跟踪或接收端信道估计误差。本报告不据此宣称任何方法最优；是否扩大性能对比范围，由后续研究目标决定。')
    table('表8 复现文件', ['文件','用途'], [
        ['verify_minimal_all_methods.m','检查Type-I/II单模式单极化退化一致性'],
        ['run_minimal_all_methods.m','执行smoke和confirm 固定全部方法与样本'],
        ['+min9/','参考方向消息 基线消息与整数解码'],
        ['analyze_minimal_all_methods.py','独立解码审计 单双流逐样本配对'],
        ['build_minimal_reports.py','由同一数据内容生成Markdown和Word'],
        ['confirm/trials.csv与paired_summary.csv','逐样本结果及同方法单双流均值'],
        ['complete_audit.json与source_and_input_SHA256.json','验证范围 数值检查及源码输入校验']
    ],[6.5,9.9])
    para('MATLAB R2024a进入第7目录，依次运行verify_minimal_all_methods、run_minimal_all_methods("smoke")和run_minimal_all_methods("confirm")；随后用含numpy、pandas、python-docx和Pillow的Python环境运行分析脚本，再运行报告构建脚本。沿用第6目录的冻结输入及算法依赖，既有+dmimo模块不作修改。')
    para('Markdown与Word采用同一组说明、公式和数据表生成。Markdown中的表格可直接检索；Word中的数据表与核心公式可编辑，转发时无需依赖本地源文件。')
    para('仓库入口 https://github.com/ttal6688/ZTE_codebokk')


def md_table(headers, rows):
    return '\n'.join(['|'+'|'.join(headers)+'|','|'+'|'.join(['---']*len(headers))+'|']+['|'+'|'.join(map(str,row))+'|' for row in rows])


def set_font(run, name='宋体', size=11, bold=False):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    for key in ['eastAsia','ascii','hAnsi']:fonts.set(qn('w:'+key),name)


def native_equation(doc, parts):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(6)
    obj=OxmlElement('m:oMath')
    def mr(text):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
    for part in parts:
        if isinstance(part,tuple) and part[0]=='sum':
            node=OxmlElement('m:sSubSup');base=OxmlElement('m:e');base.append(mr('∑'));lower=OxmlElement('m:sub');lower.append(mr(part[1]));upper=OxmlElement('m:sup');upper.append(mr(part[2]));node.extend([base,lower,upper]);obj.append(node)
        elif isinstance(part,tuple):
            node=OxmlElement('m:sSub');base=OxmlElement('m:e');base.append(mr(part[1]));sub=OxmlElement('m:sub');sub.append(mr(part[2]));node.extend([base,sub]);obj.append(node)
        else:obj.append(mr(part))
    p._p.append(obj)


def word_table(doc, title, headers, rows, widths):
    cap=doc.add_paragraph();cap.paragraph_format.keep_with_next=True;cap.paragraph_format.space_before=Pt(6);cap.paragraph_format.space_after=Pt(5);set_font(cap.add_run(title),size=10.5,bold=True)
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    if widths is None:widths=[16.4/len(headers)]*len(headers)
    for col,width in zip(tb.columns,widths):col.width=Cm(width)
    for idx,values in enumerate([headers]+rows):
        cells=tb.rows[0].cells if idx==0 else tb.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders);margin=OxmlElement('w:tcMar')
            for edge,valuepad in [('top','80'),('bottom','80'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),valuepad);e.set(qn('w:type'),'dxa');margin.append(e)
            tcpr.append(margin)
            if idx==0 or idx%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4' if idx==0 else 'F7F9FB');tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.10
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if idx==0 or (j>0 and len(headers)>=4) else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)),size=10 if len(headers)>=4 else 10.5,bold=idx==0)
        pr=tb.rows[idx]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if idx==0:pr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def write_reports():
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    # Remove built-in title/subtitle rules inherited from the default template.
    for node in doc.styles.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
    sec.footer_distance=Cm(.8)
    for name,size in [('Normal',11),('Title',20),('Heading 1',15)]:
        st=doc.styles[name];st.font.name='宋体' if name=='Normal' else '黑体';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().append(OxmlElement('w:rFonts'));st.element.rPr.rFonts.set(qn('w:eastAsia'),st.font.name)
    normal=doc.styles['Normal'].paragraph_format;normal.line_spacing=1.16;normal.space_after=Pt(5)
    title=doc.add_paragraph(style='Title');set_font(title.add_run(TITLE),'黑体',20)
    title.paragraph_format.space_after=Pt(12)
    p=doc.add_paragraph();set_font(p.add_run('全方法扩展说明与最小可行性验证    2026年10月22日'),size=10.5)
    p.paragraph_format.space_after=Pt(14)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    md=['# '+TITLE,'','全方法扩展说明与最小可行性验证  2026年10月22日','',f'[下载可转发的Word报告]({STEM}.docx)','']
    for kind,data in BLOCKS:
        if kind=='chapter':
            doc.add_page_break();p=doc.add_paragraph(data,style='Heading 1');p.paragraph_format.space_after=Pt(10);md.extend(['## '+data,''])
        elif kind=='p':
            doc.add_paragraph(data);md.extend([data,''])
        elif kind=='table':
            caption,heads,rows,widths=data;word_table(doc,caption,heads,rows,widths);md.extend(['**'+caption+'**','',md_table(heads,rows),''])
        elif kind=='eq':
            latex,parts=data;native_equation(doc,parts);md.extend(['$$',latex,'$$',''])
        elif kind=='image':
            path,caption=data;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run();r.add_picture(str(OUT/path),width=Cm(16.0))
            pr=r._r.xpath('.//wp:docPr')[0];pr.set('descr',caption)
            p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;md.extend([f'![{caption}]({path})',''])
    doc.core_properties.title=TITLE;doc.core_properties.subject='固定32位置的九方法双极化双流可行性验证';doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.save(OUT/(STEM+'.docx'));(OUT/(STEM+'.md')).write_text('\n'.join(md),encoding='utf-8')
    print('Reports written:',STEM)


if __name__=='__main__':
    pipeline();build_content();write_reports()
